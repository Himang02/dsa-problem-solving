r"""
Scrape public YouTube playlists -> title, URL, duration.

Outputs (separate from the AlgoUniversity class files):
  playlists.csv    one combined sheet, with a Playlist column
  playlists.xlsx   one worksheet per playlist (+ totals)
  playlists.docx   one section per playlist

Run:  .venv\Scripts\python.exe playlists.py
"""
import csv
import sys
from pathlib import Path

from yt_dlp import YoutubeDL

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = Path(__file__).parent

PLAYLISTS = {
    "HLD": "https://www.youtube.com/playlist?list=PL6W8uoQQ2c63W58rpNFDwdrBnq5G3EfT7",
    "LLD": "https://www.youtube.com/playlist?list=PL6W8uoQQ2c61X_9e6Net0WdYZidm7zooW",
    "Core Java": "https://www.youtube.com/playlist?list=PL6W8uoQQ2c63f469AyV78np0rbxRFppkx",
    "Springboot": "https://www.youtube.com/playlist?list=PL6W8uoQQ2c60g6_fcjDCLHSx1LBeVYqyZ",
}


def format_hms(secs):
    secs = int(secs)
    h, rem = divmod(secs, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def fetch(url):
    opts = {"quiet": True, "extract_flat": True, "skip_download": True,
            "ignoreerrors": True}
    with YoutubeDL(opts) as ydl:
        info = ydl.extract_info(url, download=False)
    rows = []
    for e in info.get("entries") or []:
        if not e or not e.get("id"):
            continue
        dur = e.get("duration")
        rows.append({
            "title": e.get("title") or "",
            "youtube_url": f"https://www.youtube.com/watch?v={e['id']}",
            "duration": format_hms(dur) if dur else "",
            "duration_seconds": int(dur) if dur else "",
        })
    return info.get("title") or "", rows


def main():
    data = {}  # name -> (playlist_title, rows)
    for name, url in PLAYLISTS.items():
        print(f"Fetching {name} ...")
        ptitle, rows = fetch(url)
        data[name] = (ptitle, rows)
        total = sum(r["duration_seconds"] for r in rows if r["duration_seconds"])
        print(f"  {len(rows)} videos, total {format_hms(total)}")

    write_csv(data)
    write_xlsx(data)
    write_docx(data)
    print("\nDone. Outputs: playlists.csv, playlists.xlsx, playlists.docx")


def write_csv(data):
    with open(HERE / "playlists.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["playlist", "title", "youtube_url",
                                          "duration", "duration_seconds"])
        w.writeheader()
        for name, (_pt, rows) in data.items():
            for r in rows:
                w.writerow({"playlist": name, **r})


def write_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    first = True
    for name, (ptitle, rows) in data.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name[:31]
        first = False
        ws.append([ptitle])
        ws["A1"].font = Font(bold=True, size=13)
        headers = ["#", "Title", "YouTube URL", "Duration", "Seconds"]
        ws.append(headers)
        for c in ws[2]:
            c.font = Font(bold=True)
        for i, r in enumerate(rows, 1):
            ws.append([i, r["title"], r["youtube_url"], r["duration"],
                       r["duration_seconds"]])
            cell = ws.cell(row=i + 2, column=3)
            cell.hyperlink = r["youtube_url"]
            cell.font = Font(color="0563C1", underline="single")
        total = sum(r["duration_seconds"] for r in rows if r["duration_seconds"])
        ws.append([])
        ws.append(["", f"TOTAL ({len(rows)} videos)", "", format_hms(total), total])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=4).font = Font(bold=True)
        for col, wd in zip("ABCDE", [5, 75, 45, 11, 9]):
            ws.column_dimensions[col].width = wd
        ws.freeze_panes = "A3"
    wb.save(HERE / "playlists.xlsx")


def write_docx(data):
    from docx import Document
    doc = Document()
    doc.add_heading("YouTube Playlists", level=0)
    for name, (ptitle, rows) in data.items():
        total = sum(r["duration_seconds"] for r in rows if r["duration_seconds"])
        doc.add_heading(f"{name} — {ptitle}", level=1)
        doc.add_paragraph(f"{len(rows)} videos · total {format_hms(total)}")
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Title", "YouTube URL", "Duration"
        for r in rows:
            cells = table.add_row().cells
            cells[0].text = r["title"]
            cells[1].text = r["youtube_url"]
            cells[2].text = r["duration"]
    doc.save(HERE / "playlists.docx")


if __name__ == "__main__":
    main()
